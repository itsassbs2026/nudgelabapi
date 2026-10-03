"""Exports (SPEC §7.2, Phase 5): files match the screens, XLSX runs as a job, everything is audited."""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import pytest
from app.config import Settings, get_settings
from app.exports import service
from app.exports.reports import REPORTS
from app.models.dashboard import Job
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import text
from sqlalchemy.orm import Session

from tests.agent_data import add_assignments, seed

SEPT = {"date_from": "2026-09-01", "date_to": "2026-09-30"}

# report → (the on-screen endpoint, its query parameters, where its rows are, the export's params)
SCREENS: dict[str, tuple[str, dict[str, Any], str, dict[str, Any]]] = {
    "sessions": ("/sessions", {}, "items", {}),
    "trainings": ("/reports/trainings", {}, "items", {}),
    "questions": ("/reports/trainings/big4/questions", {}, "questions", {"training_id": "big4"}),
    "drilldown": (
        "/reports/drilldown",
        {"level": "store", "parent": "100"},
        "rows",
        {"level": "store", "parent": "100"},
    ),
    "feedback": ("/feedback", {}, "items", {}),
    "acknowledgments": ("/acknowledgments", {}, "items", {}),
    "assignments": ("/assignments", {}, "items", {}),
    "quality": ("/quality", {}, "items", {}),
    "daily": ("/reports/overview", {}, "daily", {}),
    "cost_by_training": ("/reports/cost", {}, "by_training", {}),
}


@pytest.fixture
def data(db_session: Session) -> None:
    seed(db_session)
    add_assignments(db_session)


@pytest.fixture
def settings(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Settings:
    current = get_settings()
    monkeypatch.setattr(current, "export_dir", str(tmp_path))
    return current


def export(client: TestClient, headers: dict[str, str], report: str, fmt: str = "csv", **extra: Any) -> Any:
    body = {
        "report": report,
        "format": fmt,
        "filters": extra.pop("filters", SEPT),
        "params": extra.pop("params", {}),
    }
    return client.post("/api/v1/exports", headers=headers, json=body)


def read_csv(response: Any) -> list[list[str]]:
    assert response.status_code == 200, response.text
    assert response.content.startswith("﻿".encode())  # BOM, so Excel reads UTF-8
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def test_every_report_has_a_screen() -> None:
    assert set(SCREENS) == set(REPORTS)


@pytest.mark.parametrize("report", sorted(SCREENS))
def test_csv_matches_the_screen(
    client: TestClient, trainer_headers: dict[str, str], data: None, report: str
) -> None:
    path, query, key, params = SCREENS[report]
    screen = client.get(f"/api/v1{path}", headers=trainer_headers, params={**SEPT, **query})
    assert screen.status_code == 200, screen.text
    expected = screen.json()[key]
    rows = read_csv(export(client, trainer_headers, report, params=params))
    head, body = rows[0], rows[1:]
    assert len(head) == len(REPORTS[report].columns)
    assert len(body) == len(expected) > 0
    for line, row in zip(body, expected, strict=True):
        for cell, col in zip(line, REPORTS[report].columns, strict=True):
            value: Any = row
            for part in col.key.split("."):
                value = value[part]
            if col.kind in ("datetime",):
                continue  # converted to local time; checked below
            if value is None:
                assert cell == ""
            elif isinstance(value, bool):
                assert cell == ("yes" if value else "no")
            elif isinstance(value, list):
                assert cell == "; ".join(map(str, value))
            elif isinstance(value, int | float):
                assert float(cell) == float(value), (col.header, cell, value)
            else:
                assert cell.lstrip("'") == str(value), (col.header, cell, value)


def test_csv_times_are_local_and_named(
    client: TestClient, trainer_headers: dict[str, str], data: None
) -> None:
    rows = read_csv(export(client, trainer_headers, "sessions", params={"uid": 1001}))
    assert rows[0][1] == "Started (America/Chicago)"
    assert [r[1] for r in rows[1:]] == ["2026-09-06 10:00:00", "2026-09-05 10:00:00"]  # 15:00 UTC = 10:00 CDT


def test_formulas_never_run(
    client: TestClient, trainer_headers: dict[str, str], db_session: Session, data: None
) -> None:
    db_session.execute(
        text("UPDATE training_feedback SET comment = '=HYPERLINK(\"http://x\")' WHERE session_id = 's1'")
    )
    rows = read_csv(export(client, trainer_headers, "feedback"))
    comments = [r[6] for r in rows[1:]]
    assert '\'=HYPERLINK("http://x")' in comments


def test_csv_is_audited(
    client: TestClient, trainer_headers: dict[str, str], db_session: Session, data: None
) -> None:
    export(client, trainer_headers, "feedback", params={"min_rating": 6})
    (row,) = db_session.execute(
        text("SELECT target_id, details FROM dash_audit_log WHERE action = 'export'")
    ).all()
    details = row.details if isinstance(row.details, dict) else json.loads(row.details)
    assert row.target_id == "feedback"
    assert details["rows"] == 1 and details["format"] == "csv" and details["params"] == {"min_rating": 6}
    assert details["filters"] == SEPT


@pytest.mark.parametrize(
    ("body", "status", "code"),
    [
        ({"report": "everything", "format": "csv"}, 422, "unknown_report"),
        ({"report": "questions", "format": "csv"}, 422, "invalid_params"),  # needs a training
        (
            {"report": "drilldown", "format": "csv", "params": {"level": "market", "parent": "x"}},
            422,
            "invalid_params",
        ),
        ({"report": "daily", "format": "csv", "params": {"surprise": 1}}, 422, "invalid_params"),
        ({"report": "daily", "format": "pdf"}, 422, "validation_error"),
        ({"report": "daily", "format": "csv", "filters": {"include_bots": True}}, 403, "forbidden"),
        (
            {
                "report": "daily",
                "format": "csv",
                "filters": {"date_from": "2026-09-30", "date_to": "2026-09-01"},
            },
            422,
            "invalid_date_range",
        ),
    ],
)
def test_bad_requests(
    client: TestClient, trainer_headers: dict[str, str], body: dict[str, Any], status: int, code: str
) -> None:
    response = client.post("/api/v1/exports", headers=trainer_headers, json=body)
    assert response.status_code == status and response.json()["error"]["code"] == code


def test_row_limit(
    client: TestClient,
    trainer_headers: dict[str, str],
    monkeypatch: pytest.MonkeyPatch,
    settings: Settings,
    data: None,
) -> None:
    monkeypatch.setattr(settings, "export_max_rows", 2)
    response = export(client, trainer_headers, "sessions")
    assert response.status_code == 422 and response.json()["error"]["code"] == "export_too_large"


# -- XLSX jobs -----------------------------------------------------------------------------------------------


def test_xlsx_job_end_to_end(
    client: TestClient,
    trainer_headers: dict[str, str],
    admin_headers: dict[str, str],
    db_session: Session,
    settings: Settings,
    data: None,
) -> None:
    response = export(client, trainer_headers, "sessions", "xlsx")
    assert response.status_code == 202, response.text
    job = response.json()
    assert job["status"] == "queued" and job["report"] == "sessions"

    waiting = client.get(f"/api/v1/exports/{job['id']}/download", headers=trainer_headers)
    assert waiting.status_code == 409 and waiting.json()["error"]["code"] == "export_not_ready"

    assert service.run_export_jobs(db_session, settings) == {"done": 1, "failed": 0, "interrupted": 0}
    status = client.get(f"/api/v1/exports/{job['id']}", headers=trainer_headers).json()
    assert status["status"] == "done" and status["rows"] == 5
    assert status["filename"] == "nudgelab-sessions-2026-09-01-to-2026-09-30.xlsx"

    # Only the person who asked can see or download it.
    assert client.get(f"/api/v1/exports/{job['id']}", headers=admin_headers).status_code == 404

    download = client.get(f"/api/v1/exports/{job['id']}/download", headers=trainer_headers)
    assert download.status_code == 200
    assert "nudgelab-sessions-2026-09-01-to-2026-09-30.xlsx" in download.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(download.content))
    sheet = wb["Sessions"]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0][:3] == ("Session", "Started (America/Chicago)", "Seconds")
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"
    assert [r[0] for r in rows[1:]] == ["s5", "s4", "s3", "s2", "s1"]
    assert rows[5][1] == datetime(2026, 9, 5, 10, 0) and rows[5][15] == 0.5
    about = dict(wb["About"].iter_rows(values_only=True))
    assert about["Report"] == "Sessions" and about["Rows"] == 5

    audit = db_session.execute(text("SELECT COUNT(*) FROM dash_audit_log WHERE action = 'export'")).scalar()
    assert audit == 1


def test_old_files_are_removed(
    client: TestClient, trainer_headers: dict[str, str], db_session: Session, settings: Settings, data: None
) -> None:
    job_id = export(client, trainer_headers, "daily", "xlsx").json()["id"]
    service.run_export_jobs(db_session, settings)
    job = db_session.get(Job, job_id)
    assert job is not None and (Path(settings.export_dir) / f"{job_id}.xlsx").exists()
    job.finished_at = datetime.now() - timedelta(hours=settings.export_keep_hours + 1)
    db_session.commit()
    assert service.purge_old_exports(db_session, settings) == {"removed": 1}
    gone = client.get(f"/api/v1/exports/{job_id}/download", headers=trainer_headers)
    assert gone.status_code == 410 and gone.json()["error"]["code"] == "export_expired"
    assert client.get(f"/api/v1/exports/{job_id}", headers=trainer_headers).json()["expired"] is True


def test_failed_and_interrupted_jobs(
    client: TestClient,
    trainer_headers: dict[str, str],
    db_session: Session,
    monkeypatch: pytest.MonkeyPatch,
    settings: Settings,
    data: None,
) -> None:
    too_big = export(client, trainer_headers, "sessions", "xlsx").json()["id"]
    stuck = export(client, trainer_headers, "daily", "xlsx").json()["id"]
    job = db_session.get(Job, stuck)
    assert job is not None
    job.status, job.started_at = "running", datetime.now() - timedelta(hours=1)
    db_session.commit()

    monkeypatch.setattr(settings, "export_max_rows", 2)
    assert service.run_export_jobs(db_session, settings) == {"done": 0, "failed": 1, "interrupted": 1}
    failed = client.get(f"/api/v1/exports/{too_big}", headers=trainer_headers).json()
    assert failed["status"] == "failed" and "limited to 2 rows" in failed["error"]
    assert client.get(f"/api/v1/exports/{stuck}", headers=trainer_headers).json()["status"] == "failed"
    response = client.get(f"/api/v1/exports/{too_big}/download", headers=trainer_headers)
    assert response.status_code == 409 and response.json()["error"]["code"] == "export_failed"
