"""CSV and XLSX output for exports.

Times are shown in the user's time zone (the agent stores UTC), named in the header. Text that a spreadsheet
would run as a formula (starting with = + - @) gets a leading apostrophe: comments, quotes and transcripts are
things people said, and they must never execute in Excel.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Iterable, Iterator
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from app.exports.reports import Col, Report, value_at

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_text(text: str) -> str:
    return "'" + text if text.startswith(_FORMULA_START) else text


def headers(report: Report, timezone: str) -> list[str]:
    return [f"{c.header} ({timezone})" if c.kind == "datetime" else c.header for c in report.columns]


def local(value: datetime, tz: ZoneInfo) -> datetime:
    """Naive UTC (as stored) → naive local time, for display."""
    return value.replace(tzinfo=UTC).astimezone(tz).replace(tzinfo=None)


def typed(value: Any, col: Col, tz: ZoneInfo) -> Any:
    """The cell's value: numbers stay numbers, datetimes become local, text is made safe."""
    if value is None:
        return None
    if col.kind == "datetime" and isinstance(value, datetime):
        return local(value, tz)
    if col.kind == "date" and isinstance(value, str):
        return date.fromisoformat(value)
    if col.kind == "list" and isinstance(value, list | tuple):
        return safe_text("; ".join(str(v) for v in value))
    if isinstance(value, bool):
        return "yes" if value else "no"
    if isinstance(value, int | float):
        return value
    if isinstance(value, datetime | date):
        return value
    return safe_text(str(value))


def csv_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def csv_chunks(report: Report, rows: Iterable[dict[str, Any]], timezone: str) -> Iterator[bytes]:
    """The CSV as UTF-8 chunks, with a byte-order mark so Excel reads accents and symbols correctly."""
    tz = ZoneInfo(timezone)
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(headers(report, timezone))
    yield ("﻿" + buffer.getvalue()).encode("utf-8")
    buffer.seek(0)
    buffer.truncate()
    for i, row in enumerate(rows, start=1):
        writer.writerow([csv_text(typed(value_at(row, c.key), c, tz)) for c in report.columns])
        if i % 500 == 0:
            yield buffer.getvalue().encode("utf-8")
            buffer.seek(0)
            buffer.truncate()
    rest = buffer.getvalue()
    if rest:
        yield rest.encode("utf-8")


_NUMBER_FORMATS = {"pct": "0.0%", "float": "0.00", "datetime": "yyyy-mm-dd hh:mm", "date": "yyyy-mm-dd"}


def write_xlsx(
    path: Path, report: Report, rows: list[dict[str, Any]], timezone: str, about: dict[str, Any]
) -> None:
    """A workbook with the data sheet (bold, frozen header) and an "About" sheet describing the export."""
    tz = ZoneInfo(timezone)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(report.title[:31])
    for i, col in enumerate(report.columns):
        letter = get_column_letter(i + 1)
        wide = col.kind in ("text", "list") and col.header.lower() in (
            "comment",
            "summary",
            "statement",
            "trainee's words",
            "note",
            "text",
            "heard as",
            "issues",
        )
        ws.column_dimensions[letter].width = 60 if wide else 22 if col.kind == "datetime" else 16
    ws.freeze_panes = "A2"
    bold = Font(bold=True)
    head = []
    for text in headers(report, timezone):
        cell = WriteOnlyCell(ws, value=text)
        cell.font = bold
        head.append(cell)
    ws.append(head)
    for row in rows:
        cells = []
        for col in report.columns:
            cell = WriteOnlyCell(ws, value=typed(value_at(row, col.key), col, tz))
            if col.kind in _NUMBER_FORMATS:
                cell.number_format = _NUMBER_FORMATS[col.kind]
            cells.append(cell)
        ws.append(cells)
    info = wb.create_sheet("About")
    for key, value in about.items():
        info.append([key, value if isinstance(value, int | float | date) else safe_text(str(value or ""))])
    tmp = path.with_suffix(".part")
    wb.save(tmp)
    tmp.replace(path)
